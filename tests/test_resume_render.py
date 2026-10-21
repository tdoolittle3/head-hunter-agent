"""Markdown is canonical; DOCX and PDF are rendered from it.

So the tests check the Markdown closely and then check that the DOCX and the
PDF really came through that Markdown -- if they ever diverge, the thing the
user reads and the thing the employer opens stop agreeing.
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from pypdf import PdfReader

from head_hunter.resume_render import (
    MarkdownRenderError,
    render_docx,
    render_markdown,
    render_pdf,
)
from head_hunter.schemas import (
    Identity,
    Profile,
    Resume,
    ResumeBullet,
    ResumeSection,
)


@pytest.fixture
def profile() -> Profile:
    return Profile(
        user_id="local",
        identity=Identity(
            full_name="Jordan Reyes",
            email="jordan@example.com",
            location="Denver, CO",
            links=["linkedin.com/in/jordanreyes"],
        ),
    )


@pytest.fixture
def resume() -> Resume:
    return Resume(
        user_id="local",
        job_id="job-1",
        sections=[
            ResumeSection(
                title="Summary",
                bullets=[
                    ResumeBullet(text="Data platform engineer.", evidence_ids=["ev-1"])
                ],
            ),
            ResumeSection(
                title="Staff Engineer, Acme Foods",
                subtitle="2019-01 – 2024-11",
                bullets=[
                    ResumeBullet(text="Cut month-end close.", evidence_ids=["ev-1"]),
                    ResumeBullet(text="Mentored two engineers.", evidence_ids=["ev-2"]),
                ],
            ),
        ],
    )


def paragraphs(data: bytes) -> list[tuple[str, str]]:
    document = Document(io.BytesIO(data))
    return [(p.style.name, p.text) for p in document.paragraphs if p.text]


# Markdown ---------------------------------------------------------------


def test_markdown_has_the_expected_shape(profile: Profile, resume: Resume) -> None:
    lines = render_markdown(resume, profile).splitlines()

    assert lines[0] == "# Jordan Reyes"
    assert "## Summary" in lines
    assert "## Staff Engineer, Acme Foods" in lines
    assert "*2019-01 – 2024-11*" in lines
    assert "- Cut month-end close." in lines


def test_the_header_comes_from_the_profile_not_the_resume(
    profile: Profile, resume: Resume
) -> None:
    """Contact details are the one thing the Tailor must not be able to write."""
    markdown = render_markdown(resume, profile)

    assert "jordan@example.com" in markdown
    assert "Denver, CO" in markdown
    assert "linkedin.com/in/jordanreyes" in markdown


def test_markdown_ends_with_exactly_one_newline(
    profile: Profile, resume: Resume
) -> None:
    markdown = render_markdown(resume, profile)
    assert markdown.endswith("\n")
    assert not markdown.endswith("\n\n")


def test_a_profile_with_no_name_still_renders(resume: Resume) -> None:
    markdown = render_markdown(resume, Profile(user_id="local"))
    assert markdown.startswith("# Resume")


# DOCX -------------------------------------------------------------------


def test_docx_carries_every_line_of_the_markdown(
    profile: Profile, resume: Resume
) -> None:
    markdown = render_markdown(resume, profile)
    rendered = paragraphs(render_docx(markdown))
    texts = [text for _, text in rendered]

    assert "Jordan Reyes" in texts
    assert "Cut month-end close." in texts
    assert "Mentored two engineers." in texts
    assert any("jordan@example.com" in text for text in texts)


def test_bullets_use_the_list_style_so_an_ats_sees_a_list(
    profile: Profile, resume: Resume
) -> None:
    rendered = paragraphs(render_docx(render_markdown(resume, profile)))
    styles = {text: style for style, text in rendered}

    assert styles["Cut month-end close."] == "List Bullet"
    assert styles["Summary"].startswith("Heading")


def test_the_docx_has_nothing_an_ats_chokes_on(
    profile: Profile, resume: Resume
) -> None:
    document = Document(io.BytesIO(render_docx(render_markdown(resume, profile))))

    assert len(document.tables) == 0
    assert len(document.inline_shapes) == 0
    assert len(document.sections[0].header.paragraphs[0].text) == 0


def test_a_hand_edited_markdown_file_re_renders(profile: Profile) -> None:
    """The point of a canonical format: edit the .md and the .docx follows."""
    edited = "# Jordan Reyes\n\n## Experience\n\n- A bullet typed by hand.\n"
    texts = [text for _, text in paragraphs(render_docx(edited))]

    assert texts == ["Jordan Reyes", "Experience", "A bullet typed by hand."]


def test_inline_emphasis_survives() -> None:
    document = Document(io.BytesIO(render_docx("Led **Python** and *Airflow* work.\n")))
    runs = {run.text: run for run in document.paragraphs[0].runs}

    assert runs["Python"].bold
    assert runs["Airflow"].italic


@pytest.mark.parametrize("bad", ["| a | b |\n", "```python\nx = 1\n```\n"])
def test_unsupported_markdown_is_refused_not_dropped(bad: str) -> None:
    """Rule 6: a line silently missing from a resume is the worse failure."""
    with pytest.raises(MarkdownRenderError):
        render_docx(bad)


# PDF --------------------------------------------------------------------


def pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() for page in reader.pages)


def test_pdf_carries_every_line_of_the_markdown(
    profile: Profile, resume: Resume
) -> None:
    data = render_pdf(render_markdown(resume, profile))
    text = pdf_text(data)

    assert data[:5] == b"%PDF-"
    assert "Jordan Reyes" in text
    assert "Cut month-end close." in text
    assert "Mentored two engineers." in text
    assert "jordan@example.com" in text


def test_pdf_title_is_the_candidates_name(profile: Profile, resume: Resume) -> None:
    """The browser tab in the document viewer should say whose resume it is."""
    reader = PdfReader(io.BytesIO(render_pdf(render_markdown(resume, profile))))
    assert reader.metadata.title == "Jordan Reyes"


def test_pdf_keeps_characters_the_built_in_fonts_cannot_spell() -> None:
    """Names and evidence quotes are not ASCII; the embedded font must carry them."""
    markdown = "# Zoë Núñez-Łaska\n\n## Experience\n\n- Cut “close” — by 80%.\n"
    text = pdf_text(render_pdf(markdown))

    assert "Zoë Núñez-Łaska" in text
    assert "Cut “close” — by 80%." in text


def test_a_hand_edited_markdown_file_re_renders_as_pdf() -> None:
    edited = "# Jordan Reyes\n\n## Experience\n\n- A bullet typed by hand.\n"
    text = pdf_text(render_pdf(edited))

    assert "Jordan Reyes" in text
    assert "Experience" in text
    assert "A bullet typed by hand." in text


@pytest.mark.parametrize("bad", ["| a | b |\n", "```python\nx = 1\n```\n"])
def test_unsupported_markdown_is_refused_by_the_pdf_too(bad: str) -> None:
    with pytest.raises(MarkdownRenderError):
        render_pdf(bad)
