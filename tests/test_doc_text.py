"""Turning uploaded files into text, and refusing the ones that are not resumes.

Files arrive from the internet, so the interesting cases are the hostile ones: a
program renamed ``.pdf``, a zip that expands to gigabytes, a PDF with no text, a
password-protected one. Each must come back as a plain-language refusal, never a
stack trace and never a half-read file.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter

from head_hunter import doc_text
from head_hunter.doc_text import (
    ExtractionError,
    display_name,
    extract_text,
    is_excerpt,
    normalize_for_match,
)

RESUME_LINES = [
    "Jane Doe - Staff Engineer",
    "Acme Foods 2019-2024",
    "Rewrote the billing pipeline and cut month-end close from nine days to two.",
]


def make_pdf(lines: list[str]) -> bytes:
    """Build a tiny valid single-page PDF whose text layer is ``lines``."""

    def esc(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    stream = "BT /F1 11 Tf 50 750 Td 14 TL " + " ".join(
        f"({esc(line)}) Tj T*" for line in lines
    )
    stream += " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n"
    ).encode()
    return out


def make_docx(
    paragraphs: list[str], table: list[list[str]] | None = None, header: str = ""
) -> bytes:
    document = Document()
    if header:
        document.sections[0].header.paragraphs[0].text = header
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# What is accepted --------------------------------------------------------


def test_a_pdf_resume_is_read() -> None:
    result = extract_text("jane.pdf", make_pdf(RESUME_LINES))

    assert "billing pipeline" in result.text
    assert result.pages == 1


def test_a_word_resume_is_read_including_tables_and_the_header() -> None:
    data = make_docx(
        ["Summary of a long and varied career in data engineering."],
        table=[["Company", "Acme Foods"], ["Title", "Staff Engineer"]],
        header="Jane Doe  jane@example.com",
    )

    text = extract_text("jane.docx", data).text

    assert "jane@example.com" in text
    assert "Acme Foods" in text and "Staff Engineer" in text
    assert text.index("jane@example.com") < text.index("Summary")


@pytest.mark.parametrize("name", ["jane.txt", "jane.md", "JANE.TXT"])
def test_plain_text_resumes_are_read(name: str) -> None:
    data = "\n".join(RESUME_LINES).encode("utf-8")
    assert "month-end close" in extract_text(name, data).text


def test_a_windows_encoded_text_file_is_read() -> None:
    data = "Café manager at Acme, 2019 to 2024, running the lunch rush.".encode(
        "cp1252"
    )
    assert "Café" in extract_text("jane.txt", data).text


def test_runaway_blank_lines_and_control_characters_are_tidied() -> None:
    data = b"Jane Doe, Staff Engineer\n\n\n\n\nAcme Foods\x07 2019 to 2024 here\n"
    text = extract_text("jane.txt", data).text

    assert "\n\n\n" not in text
    assert "\x07" not in text


# What is refused ---------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "message"),
    [
        ("jane.doc", r"\.docx or PDF"),
        ("jane.exe", "Upload a PDF"),
        ("jane", "Upload a PDF"),
        ("jane.pdf.exe", "Upload a PDF"),
    ],
)
def test_unsupported_types_are_refused_with_what_to_do(name: str, message: str) -> None:
    with pytest.raises(ExtractionError, match=message):
        extract_text(name, b"some content that is long enough to be a resume")


def test_an_empty_file_is_refused() -> None:
    with pytest.raises(ExtractionError, match="empty"):
        extract_text("jane.pdf", b"")


def test_an_oversized_file_is_refused_before_it_is_parsed() -> None:
    with pytest.raises(ExtractionError, match="larger than 5 MB"):
        extract_text("jane.txt", b"x" * (doc_text.MAX_FILE_BYTES + 1))


@pytest.mark.parametrize("name", ["jane.pdf", "jane.docx"])
def test_a_program_renamed_to_look_like_a_resume_is_refused(name: str) -> None:
    with pytest.raises(ExtractionError, match="despite its name"):
        extract_text(name, b"MZ\x90\x00\x03" + b"\x00" * 200)


def test_a_binary_file_named_txt_is_refused() -> None:
    with pytest.raises(ExtractionError, match="text file"):
        extract_text("jane.txt", b"MZ\x90\x00\x03\x00\x00" + b"A" * 100)


def test_a_zip_that_is_not_a_word_document_is_refused() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("notes.txt", "hello " * 20)

    with pytest.raises(ExtractionError, match="not a Word document"):
        extract_text("jane.docx", buffer.getvalue())


def test_a_zip_bomb_posing_as_a_word_document_is_refused() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (30 * 1024 * 1024))
    assert len(buffer.getvalue()) < doc_text.MAX_FILE_BYTES  # passes the size cap

    with pytest.raises(ExtractionError, match="unreasonable size"):
        extract_text("jane.docx", buffer.getvalue())


def test_a_scanned_pdf_with_no_text_says_why_and_what_to_do() -> None:
    writer = PdfWriter()
    writer.add_blank_page(612, 792)
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(ExtractionError, match="scan.*paste the text"):
        extract_text("scan.pdf", buffer.getvalue())


def test_a_password_protected_pdf_is_refused() -> None:
    writer = PdfWriter()
    writer.add_blank_page(612, 792)
    writer.encrypt("secret")
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(ExtractionError, match="password"):
        extract_text("locked.pdf", buffer.getvalue())


def test_a_pdf_with_too_many_pages_is_refused() -> None:
    writer = PdfWriter()
    for _ in range(doc_text.MAX_PAGES + 1):
        writer.add_blank_page(612, 792)
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(ExtractionError, match="pages"):
        extract_text("book.pdf", buffer.getvalue())


def test_a_damaged_pdf_is_a_refusal_not_a_stack_trace() -> None:
    with pytest.raises(ExtractionError, match="could not be read"):
        extract_text("jane.pdf", b"%PDF-1.4\n" + b"\x00garbage\xff" * 50)


def test_a_file_that_is_nearly_empty_is_refused() -> None:
    with pytest.raises(ExtractionError, match="No readable text"):
        extract_text("jane.txt", b"hi")


def test_far_too_much_text_is_refused() -> None:
    data = ("word " * 30_000).encode()  # 150,000 characters, under the byte cap
    with pytest.raises(ExtractionError, match="far more text"):
        extract_text("novel.txt", data)


# Names -------------------------------------------------------------------


@pytest.mark.parametrize(
    ("raw", "shown"),
    [
        ("jane.pdf", "jane.pdf"),
        ("../../etc/passwd", "passwd"),
        ("C:\\Users\\jane\\resume.docx", "resume.docx"),
        ("re\x00su\x07me.pdf", "resume.pdf"),
        ("", "resume"),
        (None, "resume"),
        ("a" * 500 + ".pdf", "a" * 120),
    ],
)
def test_filenames_are_made_safe_to_show_and_store(raw: str | None, shown: str) -> None:
    assert display_name(raw) == shown


# Quoting -----------------------------------------------------------------

BULLET = (
    "• Rewrote the billing pipeline and cut month-end close\n"
    "  from nine days to two, saving ﬁnance 40 hours a week."
)


@pytest.mark.parametrize(
    "quote",
    [
        "Rewrote the billing pipeline and cut month-end close from nine days to two",
        "rewrote THE billing   pipeline",
        "cut month-end close from nine days to two, saving finance 40 hours",
        "• Rewrote the billing pipeline",
    ],
)
def test_faithful_quotes_survive_layout_differences(quote: str) -> None:
    assert is_excerpt(quote, BULLET)


@pytest.mark.parametrize(
    "quote",
    [
        "cut month-end close from nine days to one",
        "saving finance 400 hours a week",
        "Led a team of twelve engineers to rewrite the billing pipeline",
        "Rewrote the billing pipeline and the invoicing system",
        "",
        "   ",
    ],
)
def test_a_quote_that_changes_a_word_or_number_is_not_an_excerpt(
    quote: str,
) -> None:
    assert not is_excerpt(quote, BULLET)


def test_quote_style_dash_style_and_bullets_do_not_matter() -> None:
    resume = "She said “it’s done” – then shipped it • on time"
    assert is_excerpt('She said "it\'s done" - then shipped it on time', resume)


def test_normalising_keeps_every_word_and_number() -> None:
    assert normalize_for_match("  Cut 9  days\u2014to 2 ") == "cut 9 days-to 2"


def test_a_quote_survives_pdf_line_wrapping_and_hyphen_breaks() -> None:
    from head_hunter.doc_text import closest_passage, is_excerpt

    text = (
        "Built LLM-\nassisted invoice extraction into Quick\nBooks and a docs-\nsite."
    )
    assert is_excerpt("Built LLM-assisted invoice extraction into QuickBooks", text)
    assert not is_excerpt("Built invoice extraction into SAP", text)
    assert "LLM" in closest_passage("Built LLM-assisted invoice extraction", text)
