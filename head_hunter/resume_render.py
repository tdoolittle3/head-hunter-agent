"""Rendering a validated resume to Markdown and then to DOCX or PDF.

AGENTS.md: "Markdown is the canonical resume format; DOCX is rendered from it."
So there is exactly one path, and it runs in that order. :func:`render_markdown`
turns the stored record into text a person can read and edit; :func:`render_docx`
and :func:`render_pdf` read that same text back and lay it out. Hand-edit the
``.md`` and re-render and the change survives, which is the point of having a
canonical format at all.

The Markdown grammar is deliberately tiny -- headings, bullets, paragraphs, and
inline bold or italic -- because both ends of the round trip live in this file.
Both renderers fail loudly on anything they do not understand rather than
silently dropping a line off someone's resume.

The files are plain on purpose. Applicant tracking systems parse a single column
of styled paragraphs reliably and tables, text boxes, headers, and footers
badly, so none of those appear here. The PDF embeds Carlito
(``head_hunter/fonts/``), Calibri's metric twin, so the two formats look the
same and any name or evidence text survives -- the built-in PDF fonts cannot
spell half the world's names.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from fpdf import FPDF
from fpdf.enums import XPos, YPos

from head_hunter.schemas import Profile, Resume

CONTACT_SEPARATOR = " · "

_BULLET_RE = re.compile(r"^[-*]\s+(.*)$")
_HEADING_RE = re.compile(r"^(#{1,3})\s+(.*)$")
_INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")

BODY_FONT = "Calibri"
BODY_POINTS = 10.5


def _contact_line(profile: Profile) -> str:
    """Build the one-line contact row under the name."""
    identity = profile.identity
    parts = [identity.location, identity.email, identity.phone, *identity.links]
    return CONTACT_SEPARATOR.join(part for part in parts if part)


def render_markdown(resume: Resume, profile: Profile) -> str:
    """Render a resume as Markdown, the canonical format.

    The header comes from the profile's identity rather than from the resume,
    so a name or phone number can never be something the Resume Tailor made up.

    Args:
        resume: The resume to render. Its validation state is not consulted --
            callers decide whether an unvalidated draft may be rendered.
        profile: The profile the resume was written from.

    Returns:
        Markdown text ending in a single newline.
    """
    lines: list[str] = [f"# {profile.identity.full_name or 'Resume'}", ""]

    contact = _contact_line(profile)
    if contact:
        lines += [contact, ""]
    if profile.identity.headline:
        lines += [profile.identity.headline, ""]

    for section in resume.sections:
        lines.append(f"## {section.title}")
        if section.subtitle:
            lines += ["", f"*{section.subtitle}*"]
        lines.append("")
        for bullet in section.bullets:
            lines.append(f"- {bullet.text}")
        lines.append("")

    while lines and not lines[-1]:
        lines.pop()

    return "\n".join(lines) + "\n"


class MarkdownRenderError(ValueError):
    """The Markdown handed to a renderer used something unsupported.

    Raised rather than skipped: a line quietly missing from a resume is worse
    than a rendering that refuses to run (AGENTS.md rule 6).
    """


def _reject_unsupported(number: int, line: str) -> None:
    """Refuse the Markdown constructs no resume renderer here understands."""
    if line.startswith("|") or line.startswith("```"):
        raise MarkdownRenderError(
            f"Line {number} uses Markdown this renderer does not support "
            f"({line[:40]!r}). Resumes here are headings, bullets, and "
            "paragraphs only -- tables and code blocks do not survive an "
            "applicant tracking system."
        )


def _write_runs(paragraph, text: str) -> None:
    """Add ``text`` to a paragraph, honouring inline ``**bold**`` and ``*italic*``."""
    for piece in _INLINE_RE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*"):
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)


def render_docx(markdown: str) -> bytes:
    """Render canonical Markdown to an ATS-friendly DOCX.

    Args:
        markdown: Output of :func:`render_markdown`, possibly hand-edited.

    Returns:
        The ``.docx`` file as bytes, ready to write to storage.

    Raises:
        MarkdownRenderError: If the text contains Markdown this renderer does
            not support, such as a table or a fenced code block.
    """
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_POINTS)

    # Contact details sit above the first section heading and are centred with
    # the name. Everything after it is body text, left aligned.
    in_header = True

    for number, raw in enumerate(markdown.splitlines(), start=1):
        line = raw.strip()
        if not line:
            continue

        _reject_unsupported(number, line)

        heading = _HEADING_RE.match(line)
        if heading:
            level = len(heading.group(1))
            if level == 1:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(heading.group(2))
                run.bold = True
                run.font.size = Pt(18)
            else:
                in_header = False
                paragraph = document.add_heading(level=min(level, 3))
                paragraph.text = ""
                _write_runs(paragraph, heading.group(2))
            continue

        bullet = _BULLET_RE.match(line)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            _write_runs(paragraph, bullet.group(1))
            continue

        paragraph = document.add_paragraph()
        if in_header:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _write_runs(paragraph, line)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


FONT_DIR = Path(__file__).parent / "fonts"
PDF_FONT = "Carlito"
_FONT_FILES = {
    "": "Carlito-Regular.ttf",
    "B": "Carlito-Bold.ttf",
    "I": "Carlito-Italic.ttf",
    "BI": "Carlito-BoldItalic.ttf",
}

PDF_MARGIN_PT = 54  # three-quarters of an inch, a normal resume margin
PDF_BODY_POINTS = BODY_POINTS
PDF_LINE_PT = 14
PDF_BULLET_INDENT_PT = 14
_PDF_HEADING_POINTS = {1: 18.0, 2: 13.0, 3: 11.5}
_PDF_SPACE_BEFORE_PT = {1: 0.0, 2: 14.0, 3: 10.0}


def _pdf_document() -> FPDF:
    """A US-letter page with the embedded font ready in all four styles."""
    pdf = FPDF(unit="pt", format="letter")
    pdf.set_creator("Head Hunter")
    pdf.set_margins(PDF_MARGIN_PT, PDF_MARGIN_PT)
    pdf.set_auto_page_break(auto=True, margin=PDF_MARGIN_PT)
    for style, filename in _FONT_FILES.items():
        pdf.add_font(PDF_FONT, style=style, fname=str(FONT_DIR / filename))
    return pdf


def _centred_line(pdf: FPDF, text: str, size: float, style: str = "") -> None:
    """One centred line, the cursor left back at the margin on the next row."""
    pdf.set_font(PDF_FONT, style, size)
    pdf.multi_cell(0, size * 1.3, text, align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)


def _write_pdf_runs(pdf: FPDF, text: str, size: float, base_style: str = "") -> None:
    """Flow ``text`` from the cursor, honouring ``**bold**`` and ``*italic*``."""
    for piece in _INLINE_RE.split(text):
        if not piece:
            continue
        style = base_style
        if piece.startswith("**") and piece.endswith("**"):
            piece, style = piece[2:-2], base_style + "B"
        elif piece.startswith("*") and piece.endswith("*"):
            piece, style = piece[1:-1], base_style + "I"
        pdf.set_font(PDF_FONT, "".join(sorted(set(style))), size)
        pdf.write(PDF_LINE_PT, piece)
    pdf.ln(PDF_LINE_PT)


def render_pdf(markdown: str) -> bytes:
    """Render canonical Markdown to a one-column, ATS-friendly PDF.

    The same layout as :func:`render_docx`: a centred name and contact block,
    then left-aligned headings, bullets, and paragraphs.

    Args:
        markdown: Output of :func:`render_markdown`, possibly hand-edited.

    Returns:
        The ``.pdf`` file as bytes, ready to write to storage.

    Raises:
        MarkdownRenderError: If the text contains Markdown this renderer does
            not support, such as a table or a fenced code block.
    """
    pdf = _pdf_document()
    pdf.add_page()

    in_header = True

    for number, raw in enumerate(markdown.splitlines(), start=1):
        line = raw.strip()
        if not line:
            continue

        _reject_unsupported(number, line)

        heading = _HEADING_RE.match(line)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                pdf.set_title(text)
                _centred_line(pdf, text, _PDF_HEADING_POINTS[1], "B")
                continue
            in_header = False
            pdf.ln(_PDF_SPACE_BEFORE_PT[min(level, 3)])
            _write_pdf_runs(pdf, text, _PDF_HEADING_POINTS[min(level, 3)], "B")
            continue

        bullet = _BULLET_RE.match(line)
        if bullet:
            margin = pdf.l_margin
            pdf.set_font(PDF_FONT, "", PDF_BODY_POINTS)
            pdf.cell(PDF_BULLET_INDENT_PT, PDF_LINE_PT, "•")
            # Wrapped lines indent to hang under the text, not the bullet.
            pdf.set_left_margin(margin + PDF_BULLET_INDENT_PT)
            _write_pdf_runs(pdf, bullet.group(1), PDF_BODY_POINTS)
            pdf.set_left_margin(margin)
            pdf.set_x(margin)
            continue

        if in_header:
            # Centred flow cannot carry per-run styling, and the header lines
            # render_markdown writes (contact row, headline) never use it.
            _centred_line(pdf, line, PDF_BODY_POINTS)
            continue
        _write_pdf_runs(pdf, line, PDF_BODY_POINTS)

    return bytes(pdf.output())
